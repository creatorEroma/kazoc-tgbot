# -*- coding: utf-8 -*-
import asyncio
import html
import logging
import os
from datetime import date, datetime, timedelta

import asyncpg
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from aiogram import Bot, Dispatcher, F
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (CallbackQuery, FSInputFile, InlineKeyboardButton,
                           InlineKeyboardMarkup, Message, WebAppInfo)

# ───────────────────────────── КОНФИГ ─────────────────────────────
BOT_TOKEN = "8907031122:AAGms-t_ndqjcYhU9OlOIf-rECG7KWrzZw8"
DATABASE_URL = "postgresql://postgres.aijnkvwemasecsptnehy:nNarkazOc2017!@aws-1-eu-central-1.pooler.supabase.com:6543/postgres"
ADMIN_PASSWORD = "nNarkazOc2017!"
HOME_NAME = "KazOC"
REPORTS_DIR = "reports"
ADMIN_IDS = {1048016268, 8062167787}

os.makedirs(REPORTS_DIR, exist_ok=True)
logging.basicConfig(level=logging.INFO)

pool = None  # Глобальный пул БД

# Статусы
ST_ACCEPTED = "Поступило"
ST_WORK = "В работе"
ST_REOPEN = "Вернулся в работу"
ST_WIN = "Успешно"
ST_LOSS = "Провалено"
STATUS_FLOW = [ST_ACCEPTED, ST_WORK, ST_REOPEN, ST_WIN]

PERIODS = {
    "today": ("Сегодня", 0),
    "w": ("Неделя", 6),
    "m": ("Месяц", 29),
    "m3": ("3 месяца", 89),
    "m6": ("6 месяцев", 179),
    "m12": ("12 месяцев", 364),
    "all": ("Всё время", None),
}

DEAL_FIELDS = [
    ("city", "Город", "text"), ("work_date", "Дата", "date"), ("deadline", "Срок/время", "text"),
    ("hours", "Часы работы", "num"), ("position", "Позиция", "text"), 
    ("planned_count", "Количество (план)", "int"), ("work_type", "Тип работы", "text"),
    ("rate", "Ставка", "num"), ("client_pay", "Платят за заявку (клиент)", "num"),
    ("worker_pay", "Платят каждому работнику", "num"), ("company_get", "Получает компания", "num"),
    ("worker_names", "ФИО работников (смена)", "text"), ("final_workers", "Факт вышло работников", "int"),
    ("final_amount", "Итоговая сумма оплаты", "num"),
]

# ──────────────────────────── БАЗА ДАННЫХ ─────────────────────────
async def init_db():
    async with pool.acquire() as conn:
        await conn.execute("""CREATE TABLE IF NOT EXISTS managers(
            id SERIAL PRIMARY KEY, telegram_id BIGINT UNIQUE NOT NULL,
            name TEXT, username TEXT, role TEXT DEFAULT 'manager',
            active INTEGER DEFAULT 1, created_at TEXT)""")
        await conn.execute("""CREATE TABLE IF NOT EXISTS companies(
            id SERIAL PRIMARY KEY, manager_id INTEGER NOT NULL,
            name TEXT NOT NULL, phone TEXT, address TEXT, contact TEXT, notes TEXT, bitrix_id INTEGER, created_at TEXT)""")
        await conn.execute("""CREATE TABLE IF NOT EXISTS orders(
            id SERIAL PRIMARY KEY, company_id INTEGER NOT NULL, manager_id INTEGER NOT NULL,
            period_type TEXT DEFAULT 'current', city TEXT, work_date TEXT, deadline TEXT, hours REAL DEFAULT 0,
            position TEXT, planned_count INTEGER DEFAULT 0, work_type TEXT, rate REAL DEFAULT 0, 
            client_pay REAL DEFAULT 0, worker_pay REAL DEFAULT 0, company_get REAL DEFAULT 0,
            worker_names TEXT, status TEXT DEFAULT 'Принято', final_workers INTEGER DEFAULT 0, 
            final_amount REAL DEFAULT 0, notes TEXT, bitrix_id INTEGER, created_at TEXT)""")
        try:
            await conn.execute("ALTER TABLE companies ADD COLUMN IF NOT EXISTS bitrix_id INTEGER")
        except Exception:
            pass
        try:
            await conn.execute("ALTER TABLE orders ADD COLUMN IF NOT EXISTS bitrix_id INTEGER")
        except Exception:
            pass
        try:
            await conn.execute("UPDATE orders SET status='Поступило' WHERE status='Принято'")
            await conn.execute("UPDATE orders SET status='Успешно' WHERE status='Удачная сделка'")
            await conn.execute("UPDATE orders SET status='Провалено' WHERE status='Неудачная сделка'")
        except Exception:
            pass
        
        for tid in ADMIN_IDS:
            await conn.execute("""INSERT INTO managers(telegram_id, name, role, active, created_at)
                                  VALUES($1, 'Руководитель', 'admin', 1, $2)
                                  ON CONFLICT (telegram_id) DO UPDATE SET role='admin', active=1""",
                               tid, datetime.now().isoformat())

async def get_user(tid):
    async with pool.acquire() as conn:
        return await conn.fetchrow("SELECT * FROM managers WHERE telegram_id=$1", tid)

def is_allowed(u): return bool(u) and u["active"] == 1
def is_admin(u): return bool(u) and u["role"] == "admin" and u["active"] == 1

# ──────────────────────────── УТИЛИТЫ ─────────────────────────────
def esc(s): return html.escape(str(s)) if s not in (None, "") else "—"

def parse_date(text):
    t = (text or "").strip().lower()
    if t in ("", "сегодня", "today", "-"): return date.today().isoformat()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d.%m", "%Y-%m-%d"):
        try:
            d = datetime.strptime(t, fmt).date()
            if fmt == "%d.%m": d = d.replace(year=date.today().year)
            return d.isoformat()
        except ValueError: continue
    return None

def parse_num(text):
    raw = (text or "").replace(",", ".").strip()
    if raw in ("-", ""): return 0.0
    try: return float("".join(ch for ch in raw if ch.isdigit() or ch == ".") or "0")
    except ValueError: return 0.0

def parse_int(text):
    digits = "".join(ch for ch in (text or "") if ch.isdigit())
    return int(digits) if digits else 0

def fmt_date(iso):
    try: return datetime.fromisoformat(iso).strftime("%d.%m.%Y")
    except Exception: return iso or "—"

def money(v): return f"{float(v or 0):,.0f}".replace(",", " ")
def status_emoji(st): return {ST_ACCEPTED: "🟡", ST_WORK: "🔵", ST_REOPEN: "↩️", ST_WIN: "✅", ST_LOSS: "❌"}.get(st, "⚪")

def status_bar(st):
    if st == ST_LOSS: return "❌ <b>Провалено</b>"
    parts = []
    cur = STATUS_FLOW.index(st) if st in STATUS_FLOW else 0
    for i, s in enumerate(STATUS_FLOW):
        mark = "🟢" if i < cur else ("🔵" if i == cur else "⚪")
        label = f"<b>{s}</b>" if i == cur else s
        parts.append(f"{mark} {label}")
    return " → ".join(parts)

# ──────────────────────────── СОСТОЯНИЯ ───────────────────────────
class CompanyForm(StatesGroup): name = State(); phone = State(); address = State(); contact = State(); notes = State()
class DealForm(StatesGroup): city = State(); work_date = State(); deadline = State(); hours = State(); position = State(); planned_count = State(); work_type = State(); rate = State(); client_pay = State(); worker_pay = State(); company_get = State(); worker_names = State()
class CompleteForm(StatesGroup): final_workers = State(); final_amount = State()
class EditField(StatesGroup): value = State()
class AdminAuth(StatesGroup): password = State()
class ManagerForm(StatesGroup): telegram_id = State(); name = State()
class RenameManager(StatesGroup): name = State()
class ReportDay(StatesGroup): day = State()

# ──────────────────────────── КЛАВИАТУРЫ ──────────────────────────
def kb(rows): return InlineKeyboardMarkup(inline_keyboard=rows)
def btn(text, data): return InlineKeyboardButton(text=text, callback_data=data)

def main_menu(u):
    web_url = os.environ.get("WEB_APP_URL", "https://kazoc-tgbot.vercel.app")
    rows = [
        [InlineKeyboardButton(text="📱 Открыть Web App", web_app=WebAppInfo(url=web_url))],
        [btn("📋 Компании-клиенты", "co:list:0")],
        [btn("➕ Новая компания", "co:new")],
        [btn("📊 Отчёты", "rep:menu")]
    ]
    if is_admin(u): rows.append([btn("🔐 Админка", "adm:auth")])
    return kb(rows)

def period_kb(scope, sid, back):
    rows, line = [], []
    for code, (label, _) in PERIODS.items():
        line.append(btn(label, f"rep:go:{scope}:{sid}:{code}"))
        if len(line) == 2: rows.append(line); line = []
    if line: rows.append(line)
    rows.append([btn("📅 Конкретный день", f"rep:day:{scope}:{sid}")])
    rows.append([btn("⬅️ Назад", back)])
    return kb(rows)

# ──────────────────────────── ДОСТУП ──────────────────────────────
async def companies_for(u):
    async with pool.acquire() as conn:
        if is_admin(u): return await conn.fetch("SELECT * FROM companies ORDER BY name")
        return await conn.fetch("SELECT * FROM companies WHERE manager_id=$1 ORDER BY name", u["id"])

async def get_company(cid):
    async with pool.acquire() as conn: return await conn.fetchrow("SELECT * FROM companies WHERE id=$1", cid)

async def get_deal(did):
    async with pool.acquire() as conn: return await conn.fetchrow("SELECT * FROM orders WHERE id=$1", did)

# ───────────────────────────── BITRIX24 INTEGRATION ───────────────────
BITRIX_URL = os.environ.get("BITRIX_URL", "https://b24-zjchkj.bitrix24.kz/rest/1/jaf58v0tf4ioi6kx/")

async def sync_company_to_bitrix(company_id: int) -> int:
    import httpx
    try:
        company = await get_company(company_id)
        if not company: return None
        
        payload = {
            "fields": {
                "TITLE": company["name"],
                "PHONE": [{"VALUE": company["phone"], "VALUE_TYPE": "WORK"}] if company["phone"] else [],
                "ADDRESS": company["address"] if company["address"] else "",
                "COMMENTS": f"Контактное лицо: {esc(company['contact'])}<br/>Заметки: {esc(company['notes'])}"
            }
        }
        
        async with httpx.AsyncClient() as client:
            if company.get("bitrix_id"):
                await client.post(f"{BITRIX_URL}crm.company.update.json", json={
                    "id": company["bitrix_id"],
                    "ID": company["bitrix_id"],
                    "fields": payload["fields"],
                    "FIELDS": payload["fields"]
                })
                return company["bitrix_id"]
            else:
                res = await client.post(f"{BITRIX_URL}crm.company.add.json", json={
                    "fields": payload["fields"],
                    "FIELDS": payload["fields"]
                })
                if res.status_code == 200:
                    rdata = res.json()
                    if "result" in rdata:
                        bid = int(rdata["result"])
                        async with pool.acquire() as conn:
                            await conn.execute("UPDATE companies SET bitrix_id=$1 WHERE id=$2", bid, company_id)
                        return bid
    except Exception as e:
        logging.error(f"Bitrix24 company sync error: {e}")
    return None

async def sync_deal_to_bitrix(deal_id: int):
    import httpx
    try:
        deal = await get_deal(deal_id)
        if not deal: return
        company = await get_company(deal["company_id"])
        co_name = company["name"] if company else "Неизвестно"
        
        co_bitrix_id = await sync_company_to_bitrix(deal["company_id"])
        
        async with pool.acquire() as conn:
            mgr = await conn.fetchrow("SELECT name, role FROM managers WHERE id=$1", deal["manager_id"])
        if mgr:
            role_label = "Руководитель" if mgr["role"] == "admin" else "Менеджер"
            mgr_name = f"{mgr['name']} ({role_label})"
        else:
            mgr_name = "Неизвестно"
        
        stage_map = {ST_ACCEPTED: "NEW", ST_WORK: "PREPARATION", ST_REOPEN: "PREPAYMENT_INVOICE", ST_WIN: "WON", ST_LOSS: "LOSE"}
        stage_id = stage_map.get(deal["status"], "NEW")
        
        comments = f"""
        <b>Канал:</b> KazOC CRM Bot<br/>
        <b>Компания:</b> {esc(co_name)}<br/>
        <b>Менеджер:</b> {esc(mgr_name)}<br/>
        <b>Город:</b> {esc(deal['city'])}<br/>
        <b>Дата смены:</b> {fmt_date(deal['work_date'])}<br/>
        <b>Время/Срок:</b> {esc(deal['deadline'])}<br/>
        <b>Часы работы:</b> {deal['hours'] or 0}ч<br/>
        <b>Позиция:</b> {esc(deal['position'])}<br/>
        <b>Тип работы:</b> {esc(deal['work_type'])}<br/>
        <b>Количество (план):</b> {deal['planned_count']} чел<br/>
        <b>Ставка:</b> {deal['rate']} ₸<br/>
        <b>Клиент платит:</b> {deal['client_pay']} ₸<br/>
        <b>Выплата рабочему:</b> {deal['worker_pay']} ₸<br/>
        <b>Компания получает:</b> {deal['company_get']} ₸<br/>
        <b>Рабочие:</b> {esc(deal['worker_names'])}<br/>
        """
        if deal["status"] == ST_WIN:
            comments += f"""
            <br/><b>🏁 Результаты закрытия:</b><br/>
            <b>Фактически вышло рабочих:</b> {deal['final_workers']} чел<br/>
            <b>Фактическая сумма оплаты:</b> {deal['final_amount']} ₸<br/>
            """
        
        payload = {
            "fields": {
                "TITLE": f"[KazOC] {esc(deal['position'])} - {esc(co_name)} | {esc(mgr_name)}",
                "STAGE_ID": stage_id,
                "OPPORTUNITY": deal["final_amount"] if deal["status"] == ST_WIN else deal["client_pay"],
                "CURRENCY_ID": "KZT",
                "COMMENTS": comments
            }
        }
        if co_bitrix_id:
            payload["fields"]["COMPANY_ID"] = co_bitrix_id
            
        async with httpx.AsyncClient() as client:
            if deal.get("bitrix_id"):
                await client.post(f"{BITRIX_URL}crm.deal.update.json", json={
                    "id": deal["bitrix_id"],
                    "ID": deal["bitrix_id"],
                    "fields": payload["fields"],
                    "FIELDS": payload["fields"]
                })
            else:
                res = await client.post(f"{BITRIX_URL}crm.deal.add.json", json={
                    "fields": payload["fields"],
                    "FIELDS": payload["fields"]
                })
                if res.status_code == 200:
                    rdata = res.json()
                    if "result" in rdata:
                        bid = int(rdata["result"])
                        async with pool.acquire() as conn:
                            await conn.execute("UPDATE orders SET bitrix_id=$1 WHERE id=$2", bid, deal_id)
    except Exception as e:
        logging.error(f"Bitrix24 sync error: {e}")

def can_access_company(u, c): return c and (is_admin(u) or c["manager_id"] == u["id"])
def can_access_deal(u, o): return o and (is_admin(u) or o["manager_id"] == u["id"])
def can_edit_deal(u, o):
    if not o: return False
    if is_admin(u): return True
    if o["manager_id"] != u["id"]: return False
    return (o["created_at"] or "")[:10] == date.today().isoformat()

# ──────────────────────────── ОТЧЁТЫ XLSX ─────────────────────────
async def fetch_orders(scope, sid, d_from, d_to):
    query = """SELECT o.*, c.name AS company_name, m.name AS manager_name
               FROM orders o JOIN companies c ON c.id=o.company_id
               JOIN managers m ON m.id=o.manager_id WHERE 1=1"""
    args = []
    if d_from:
        args.append(d_from)
        query += f" AND o.work_date>=${len(args)}"
    if d_to:
        args.append(d_to)
        query += f" AND o.work_date<=${len(args)}"
    if scope == "comp":
        args.append(sid)
        query += f" AND o.company_id=${len(args)}"
    elif scope in ("mgr", "my"):
        args.append(sid)
        query += f" AND o.manager_id=${len(args)}"
    query += " ORDER BY o.work_date, o.id"
    async with pool.acquire() as conn:
        return await conn.fetch(query, *args)

def period_range(code):
    _, days = PERIODS[code]
    if days is None: return None, None
    return (date.today() - timedelta(days=days)).isoformat(), date.today().isoformat()

def build_excel(rows, title, period_label):
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="2F5496")
    head_font = Font(bold=True, color="FFFFFF")
    money_fmt = "#,##0"

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Сводка"
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Период: {period_label}"
    ws["A3"] = f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A4"] = f"Всего заявок: {len(rows)}"
    st_counts = {ST_ACCEPTED: 0, ST_WORK: 0, ST_WIN: 0, ST_LOSS: 0}
    for r in rows: st_counts[r["status"]] = st_counts.get(r["status"], 0) + 1
    ws["A5"] = f"Статусы — Принято: {st_counts.get(ST_ACCEPTED,0)} · В работе: {st_counts.get(ST_WORK,0)} · Удачных: {st_counts.get(ST_WIN,0)} · Неудачных: {st_counts.get(ST_LOSS,0)}"
    ws["A6"] = f"Факт вышло работников (итого): {sum(int(r['final_workers'] or 0) for r in rows)}"
    ws["A7"] = f"Итоговая сумма оплаты: {money(sum(float(r['final_amount'] or 0) for r in rows))}"

    by_co = {}
    for r in rows:
        d = by_co.setdefault(r["company_name"], {"o": 0, "plan": 0, "fact": 0, "cli": 0.0, "amt": 0.0})
        d["o"] += 1; d["plan"] += int(r["planned_count"] or 0); d["fact"] += int(r["final_workers"] or 0)
        d["cli"] += float(r["client_pay"] or 0); d["amt"] += float(r["final_amount"] or 0)

    hr = 9
    heads = ["Компания", "Заявок", "План раб.", "Факт вышло", "Платят клиенты", "Итог. сумма"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(hr, c, h); cell.fill = head_fill; cell.font = head_font; cell.border = border
    ri = hr + 1
    for name, d in sorted(by_co.items()):
        vals = [name, d["o"], d["plan"], d["fact"], round(d["cli"]), round(d["amt"])]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(ri, c, v); cell.border = border
            if c in (5, 6): cell.number_format = money_fmt
        ri += 1
    for col, w in zip("ABCDEF", (32, 9, 11, 12, 16, 14)): ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Заявки")
    cols = ["Дата", "Компания", "Город", "Менеджер", "Позиция", "Тип работы", "Период", "План кол-во", "Факт вышло", "Часы", "Ставка", "Платят за заявку", "Платят работнику", "Получает компания", "Итог. сумма", "Статус", "Срок/время", "ФИО работников", "Создано"]
    for c, h in enumerate(cols, 1):
        cell = ws2.cell(1, c, h); cell.fill = head_fill; cell.font = head_font; cell.border = border; cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, r in enumerate(rows, start=2):
        pt = "Будущий" if r["period_type"] == "future" else "Текущий"
        vals = [fmt_date(r["work_date"]), r["company_name"], r["city"] or "", r["manager_name"], r["position"] or "", r["work_type"] or "", pt, int(r["planned_count"] or 0), int(r["final_workers"] or 0), float(r["hours"] or 0), round(float(r["rate"] or 0)), round(float(r["client_pay"] or 0)), round(float(r["worker_pay"] or 0)), round(float(r["company_get"] or 0)), round(float(r["final_amount"] or 0)), r["status"] or "", r["deadline"] or "", r["worker_names"] or "", (r["created_at"] or "")[:16].replace("T", " ")]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(i, c, v); cell.border = border
            if c in (11, 12, 13, 14, 15): cell.number_format = money_fmt
    widths = [11, 24, 13, 16, 18, 18, 9, 11, 11, 7, 10, 15, 15, 16, 13, 16, 16, 30, 17]
    for c, w in enumerate(widths, 1): ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "A2"

    path = os.path.join(REPORTS_DIR, f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    wb.save(path)
    return path

# ──────────────────────────── БОТ ─────────────────────────────────
dp = Dispatcher(storage=MemoryStorage())

async def safe_edit(cq, text, markup):
    try: await cq.message.edit_text(text, reply_markup=markup)
    except Exception: await cq.message.answer(text, reply_markup=markup)

@dp.message(CommandStart())
@dp.message(Command("menu"))
async def cmd_start(m: Message, state: FSMContext):
    await state.clear()
    u = await get_user(m.from_user.id)
    if not is_allowed(u):
        await m.answer(f"⛔️ <b>Доступ запрещён.</b>\nВаш Telegram ID: <code>{m.from_user.id}</code>")
        return
    await m.answer(f"🏠 <b>{esc(HOME_NAME)}</b>\nЗдравствуйте, <b>{esc(u['name'] or 'пользователь')}</b>.\nНиже — компании, которыми вы управляете:", reply_markup=main_menu(u))

@dp.callback_query(F.data == "menu:main")
async def to_main(cq: CallbackQuery, state: FSMContext):
    await state.clear()
    u = await get_user(cq.from_user.id)
    if not is_allowed(u): await cq.answer("Доступ запрещён", show_alert=True); return
    await safe_edit(cq, f"🏠 <b>{esc(HOME_NAME)}</b>\nГлавное меню:", main_menu(u))
    await cq.answer()

@dp.callback_query(F.data == "co:new")
async def co_new(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id)
    if not is_allowed(u): await cq.answer("Доступ", show_alert=True); return
    await state.clear(); await state.set_state(CompanyForm.name)
    await cq.message.answer("🏢 <b>Новая компания.</b>\nНазвание:")
    await cq.answer()

@dp.message(CompanyForm.name)
async def co_name(m: Message, state: FSMContext): await state.update_data(name=m.text.strip()); await state.set_state(CompanyForm.phone); await m.answer("📞 Телефон (или «-»):")
@dp.message(CompanyForm.phone)
async def co_phone(m: Message, state: FSMContext): await state.update_data(phone=m.text.strip()); await state.set_state(CompanyForm.address); await m.answer("📍 Адрес (или «-»):")
@dp.message(CompanyForm.address)
async def co_addr(m: Message, state: FSMContext): await state.update_data(address=m.text.strip()); await state.set_state(CompanyForm.contact); await m.answer("👤 Контакт (или «-»):")
@dp.message(CompanyForm.contact)
async def co_contact(m: Message, state: FSMContext): await state.update_data(contact=m.text.strip()); await state.set_state(CompanyForm.notes); await m.answer("🗒 Заметки (или «-»):")

@dp.message(CompanyForm.notes)
async def co_save(m: Message, state: FSMContext):
    d = await state.get_data(); u = await get_user(m.from_user.id)
    clean = lambda v: None if v in ("-", "") else v
    async with pool.acquire() as conn:
        if d.get("edit_id"):
            await conn.execute("UPDATE companies SET name=$1, phone=$2, address=$3, contact=$4, notes=$5 WHERE id=$6", d["name"], clean(d["phone"]), clean(d["address"]), clean(d["contact"]), clean(m.text.strip()), d["edit_id"])
            cid = d["edit_id"]
        else:
            cid = await conn.fetchval("INSERT INTO companies(manager_id, name, phone, address, contact, notes, created_at) VALUES($1,$2,$3,$4,$5,$6,$7) RETURNING id", u["id"], d["name"], clean(d["phone"]), clean(d["address"]), clean(d["contact"]), clean(m.text.strip()), datetime.now().isoformat())
    asyncio.create_task(sync_company_to_bitrix(cid))
    await state.clear()
    await m.answer("✅ Компания сохранена.", reply_markup=kb([[btn("🔎 Открыть", f"co:view:{cid}")], [btn("🏠 Меню", "menu:main")]]))

@dp.callback_query(F.data.startswith("co:list:"))
async def co_list(cq: CallbackQuery, state: FSMContext):
    await state.clear(); u = await get_user(cq.from_user.id)
    if not is_allowed(u): return await cq.answer("Доступ запрещён")
    page = int(cq.data.split(":")[2]); per = 8
    items = await companies_for(u)
    if not items:
        await safe_edit(cq, f"🏠 {esc(HOME_NAME)}\nКомпаний пока нет.", kb([[btn("➕ Новая компания", "co:new")], [btn("🏠 Меню", "menu:main")]])); return await cq.answer()
    chunk = items[page * per:(page + 1) * per]
    rows = [[btn(f"🏢 {c['name']}", f"co:view:{c['id']}")] for c in chunk]
    nav = []
    if page > 0: nav.append(btn("⬅️", f"co:list:{page-1}"))
    if (page + 1) * per < len(items): nav.append(btn("➡️", f"co:list:{page+1}"))
    if nav: rows.append(nav)
    rows.append([btn("➕ Новая компания", "co:new"), btn("🏠 Меню", "menu:main")])
    await safe_edit(cq, f"🏠 <b>{esc(HOME_NAME)}</b>\nКомпании ({len(items)}):", kb(rows))
    await cq.answer()

@dp.callback_query(F.data.startswith("co:view:"))
async def co_view(cq: CallbackQuery, state: FSMContext):
    await state.clear(); u = await get_user(cq.from_user.id)
    cid = int(cq.data.split(":")[2]); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    async with pool.acquire() as conn:
        agg = await conn.fetchrow("SELECT COUNT(*) as n, COALESCE(SUM(final_workers),0) as fw, COALESCE(SUM(final_amount),0) as fa FROM orders WHERE company_id=$1", cid)
    txt = (f"🏢 <b>{esc(c['name'])}</b>\n📞 {esc(c['phone'])}  ·  📍 {esc(c['address'])}\n👤 {esc(c['contact'])}\n🗒 {esc(c['notes'])}\n\nЗаявок: <b>{agg['n']}</b> · вышло: <b>{agg['fw']}</b> · оплата: <b>{money(agg['fa'])}</b>")
    rows = [[btn("➕ Новая заявка/сделка", f"deal:new:{cid}")], [btn("🧾 Текущие", f"deal:list:{cid}:current"), btn("🔜 Будущие", f"deal:list:{cid}:future")], [btn("📊 Отчёт", f"rep:p:comp:{cid}")], [btn("✏️ Изменить", f"co:edit:{cid}"), btn("🗑 Удалить", f"co:del:{cid}")], [btn("📋 К списку", "co:list:0"), btn("🏠 Меню", "menu:main")]]
    await safe_edit(cq, txt, kb(rows)); await cq.answer()

@dp.callback_query(F.data.startswith("co:edit:"))
async def co_edit(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id); cid = int(cq.data.split(":")[2]); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    await state.clear(); await state.update_data(edit_id=cid); await state.set_state(CompanyForm.name)
    await cq.message.answer(f"✏️ Изменение «{esc(c['name'])}». Новое название:"); await cq.answer()

@dp.callback_query(F.data.startswith("co:del:"))
async def co_del(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); cid = int(cq.data.split(":")[2]); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    await safe_edit(cq, f"⚠️ Удалить «{esc(c['name'])}» и ВСЕ заявки?", kb([[btn("❌ Да, удалить", f"co:delok:{cid}")], [btn("Отмена", f"co:view:{cid}")]])); await cq.answer()

@dp.callback_query(F.data.startswith("co:delok:"))
async def co_delok(cq: CallbackQuery):
    import httpx
    u = await get_user(cq.from_user.id); cid = int(cq.data.split(":")[2]); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    if c.get("bitrix_id"):
        try:
            async with httpx.AsyncClient() as client:
                await client.post(f"{BITRIX_URL}crm.company.delete.json", json={"id": c["bitrix_id"]})
        except Exception:
            pass
    async with pool.acquire() as conn:
        await conn.execute("DELETE FROM orders WHERE company_id=$1", cid)
        await conn.execute("DELETE FROM companies WHERE id=$1", cid)
    await safe_edit(cq, "🗑 Удалено.", kb([[btn("📋 К списку", "co:list:0")]])); await cq.answer()

@dp.callback_query(F.data.startswith("deal:list:"))
async def deal_list(cq: CallbackQuery, state: FSMContext):
    await state.clear(); u = await get_user(cq.from_user.id); _, _, cid, pt = cq.data.split(":"); cid = int(cid); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    async with pool.acquire() as conn:
        ds = await conn.fetch("SELECT * FROM orders WHERE company_id=$1 AND period_type=$2 ORDER BY work_date DESC, id DESC LIMIT 30", cid, pt)
    title = "🧾 Текущие" if pt == "current" else "🔜 Будущие"
    if not ds:
        await safe_edit(cq, f"{title} — «{esc(c['name'])}»\n\nПусто.", kb([[btn("➕ Новая заявка", f"deal:new:{cid}")], [btn("⬅️ Назад", f"co:view:{cid}")]])); return await cq.answer()
    rows = [[btn(f"{status_emoji(o['status'])} {fmt_date(o['work_date'])} · {(o['position'] or o['work_type'] or 'заявка')[:18]} · {o['planned_count']}чел", f"deal:view:{o['id']}")] for o in ds]
    rows.append([btn("➕ Новая заявка", f"deal:new:{cid}"), btn("⬅️ Назад", f"co:view:{cid}")])
    await safe_edit(cq, f"{title} — «{esc(c['name'])}» ({len(ds)}):", kb(rows)); await cq.answer()

@dp.callback_query(F.data.startswith("deal:new:"))
async def deal_new(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id); cid = int(cq.data.split(":")[2]); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    await safe_edit(cq, f"📝 Новая заявка. Какой период?", kb([[btn("🧾 Текущий", f"deal:pt:{cid}:current")], [btn("🔜 Будущий", f"deal:pt:{cid}:future")], [btn("⬅️ Назад", f"co:view:{cid}")]])); await cq.answer()

@dp.callback_query(F.data.startswith("deal:pt:"))
async def deal_pt(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id); _, _, cid, pt = cq.data.split(":"); cid = int(cid); c = await get_company(cid)
    if not can_access_company(u, c): return await cq.answer("Нет доступа")
    await state.clear(); await state.update_data(company_id=cid, company_name=c["name"], period_type=pt); await state.set_state(DealForm.city)
    await cq.message.answer(f"🏢 «{esc(c['name'])}»\n🏙 Город:"); await cq.answer()

@dp.message(DealForm.city)
async def d_city(m: Message, state: FSMContext): await state.update_data(city=m.text.strip()); await state.set_state(DealForm.work_date); await m.answer("📅 Дата работ (ДД.ММ.ГГГГ или «сегодня»):")
@dp.message(DealForm.work_date)
async def d_date(m: Message, state: FSMContext):
    d = parse_date(m.text)
    if d is None: return await m.answer("Не понял дату:")
    await state.update_data(work_date=d); await state.set_state(DealForm.deadline); await m.answer("⏰ Срок/время (или «-»):")
@dp.message(DealForm.deadline)
async def d_deadline(m: Message, state: FSMContext): v = None if m.text.strip() in ("-", "") else m.text.strip(); await state.update_data(deadline=v); await state.set_state(DealForm.hours); await m.answer("🕒 Часы работы (число или «-»):")
@dp.message(DealForm.hours)
async def d_hours(m: Message, state: FSMContext): await state.update_data(hours=parse_num(m.text)); await state.set_state(DealForm.position); await m.answer("📌 Позиция:")
@dp.message(DealForm.position)
async def d_position(m: Message, state: FSMContext): await state.update_data(position=m.text.strip()); await state.set_state(DealForm.planned_count); await m.answer("👷 Количество (план):")
@dp.message(DealForm.planned_count)
async def d_count(m: Message, state: FSMContext): await state.update_data(planned_count=parse_int(m.text)); await state.set_state(DealForm.work_type); await m.answer("🔧 Тип работы:")
@dp.message(DealForm.work_type)
async def d_type(m: Message, state: FSMContext): await state.update_data(work_type=m.text.strip()); await state.set_state(DealForm.rate); await m.answer("📊 Ставка:")
@dp.message(DealForm.rate)
async def d_rate(m: Message, state: FSMContext): await state.update_data(rate=parse_num(m.text)); await state.set_state(DealForm.client_pay); await m.answer("💵 Платят за заявку (клиент):")
@dp.message(DealForm.client_pay)
async def d_clientpay(m: Message, state: FSMContext): await state.update_data(client_pay=parse_num(m.text)); await state.set_state(DealForm.worker_pay); await m.answer("👛 Платят работнику:")
@dp.message(DealForm.worker_pay)
async def d_workerpay(m: Message, state: FSMContext): await state.update_data(worker_pay=parse_num(m.text)); await state.set_state(DealForm.company_get); await m.answer("🏦 Получает компания:")
@dp.message(DealForm.company_get)
async def d_companyget(m: Message, state: FSMContext): await state.update_data(company_get=parse_num(m.text)); await state.set_state(DealForm.worker_names); await m.answer("🧑‍🤝‍🧑 ФИО работников (или «-»):")

@dp.message(DealForm.worker_names)
async def d_save(m: Message, state: FSMContext):
    d = await state.get_data(); u = await get_user(m.from_user.id); names = None if m.text.strip() in ("-", "") else m.text.strip()
    async with pool.acquire() as conn:
        did = await conn.fetchval("""INSERT INTO orders(company_id,manager_id,period_type,city,work_date,deadline,hours,position,planned_count,work_type,rate,client_pay,worker_pay,company_get,worker_names,status,created_at) VALUES($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15,$16,$17) RETURNING id""", d["company_id"], u["id"], d["period_type"], d["city"], d["work_date"], d["deadline"], d["hours"], d["position"], d["planned_count"], d["work_type"], d["rate"], d["client_pay"], d["worker_pay"], d["company_get"], names, ST_ACCEPTED, datetime.now().isoformat())
    asyncio.create_task(sync_deal_to_bitrix(did))
    await state.clear()
    text, markup = await render_deal(await get_deal(did), u)
    await m.answer("✅ Заявка создана.\n\n" + text, reply_markup=markup)

async def render_deal(o, u):
    c = await get_company(o["company_id"])
    pt = "Будущий" if o["period_type"] == "future" else "Текущий"
    txt = (f"{status_emoji(o['status'])} <b>Заявка #{o['id']}</b> · {pt}\n{status_bar(o['status'])}\n\n🏢 {esc(c['name'] if c else '—')}\n🏙 Город: {esc(o['city'])}\n📅 Дата: {fmt_date(o['work_date'])}  ⏰ {esc(o['deadline'])}  🕒 {o['hours'] or 0}ч\n📌 Позиция: {esc(o['position'])}  ·  🔧 {esc(o['work_type'])}\n👷 План: {o['planned_count'] or 0}  ·  📊 Ставка: {money(o['rate'])}\n💵 Клиент: {money(o['client_pay'])}\n👛 Работнику: {money(o['worker_pay'])}  ·  🏦 Компании: {money(o['company_get'])}\n🧑‍🤝‍🧑 ФИО: {esc(o['worker_names'])}\n")
    if o["status"] == ST_WIN: txt += (f"\n🏁 <b>Факт вышло: {o['final_workers'] or 0}</b>\n<b>Сумма оплаты: {money(o['final_amount'])}</b>\n")
    st_rows = []
    if o["status"] == ST_ACCEPTED: st_rows = [[btn("▶️ В работу", f"deal:st:{o['id']}:work"), btn("❌ Неудачная", f"deal:st:{o['id']}:loss")]]
    elif o["status"] in (ST_WORK, ST_REOPEN): st_rows = [[btn("✅ Завершить (удачно)", f"deal:st:{o['id']}:win")], [btn("❌ Неудачная", f"deal:st:{o['id']}:loss")]]
    elif o["status"] in (ST_WIN, ST_LOSS): st_rows = [[btn("↩️ Вернуть в работу", f"deal:st:{o['id']}:reopen")]]
    rows = list(st_rows)
    if can_edit_deal(u, o): rows.append([btn("✏️ Редактировать", f"deal:edit:{o['id']}"), btn("🗑 Удалить", f"deal:rm:{o['id']}")])
    rows.append([btn("⬅️ К компании", f"co:view:{o['company_id']}")])
    return txt, kb(rows)

@dp.callback_query(F.data.startswith("deal:view:"))
async def deal_view(cq: CallbackQuery, state: FSMContext):
    await state.clear(); u = await get_user(cq.from_user.id); o = await get_deal(int(cq.data.split(":")[2]))
    if not can_access_deal(u, o): return await cq.answer("Нет доступа")
    text, markup = await render_deal(o, u)
    await safe_edit(cq, text, markup); await cq.answer()

@dp.callback_query(F.data.startswith("deal:st:"))
async def deal_status(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id); _, _, did, action = cq.data.split(":"); did = int(did); o = await get_deal(did)
    if not can_access_deal(u, o): return await cq.answer("Нет доступа")
    if action == "win":
        await state.clear(); await state.update_data(deal_id=did); await state.set_state(CompleteForm.final_workers); await cq.message.answer("🏁 Сколько работников вышло?"); return await cq.answer()
    new = {"work": ST_WORK, "loss": ST_LOSS, "reopen": ST_REOPEN}[action]
    async with pool.acquire() as conn: await conn.execute("UPDATE orders SET status=$1 WHERE id=$2", new, did)
    asyncio.create_task(sync_deal_to_bitrix(did))
    await cq.answer("Статус обновлён"); text, markup = await render_deal(await get_deal(did), u); await safe_edit(cq, text, markup)

@dp.message(CompleteForm.final_workers)
async def cf_workers(m: Message, state: FSMContext): await state.update_data(final_workers=parse_int(m.text)); await state.set_state(CompleteForm.final_amount); await m.answer("💰 Итоговая сумма оплаты:")
@dp.message(CompleteForm.final_amount)
async def cf_amount(m: Message, state: FSMContext):
    d = await state.get_data(); u = await get_user(m.from_user.id)
    async with pool.acquire() as conn: await conn.execute("UPDATE orders SET status=$1, final_workers=$2, final_amount=$3 WHERE id=$4", ST_WIN, d["final_workers"], parse_num(m.text), d["deal_id"])
    asyncio.create_task(sync_deal_to_bitrix(d["deal_id"]))
    await state.clear(); text, markup = await render_deal(await get_deal(d["deal_id"]), u); await m.answer("✅ Завершено.\n\n" + text, reply_markup=markup)

@dp.callback_query(F.data.startswith("deal:edit:"))
async def deal_edit(cq: CallbackQuery, state: FSMContext):
    await state.clear(); u = await get_user(cq.from_user.id); did = int(cq.data.split(":")[2]); o = await get_deal(did)
    if not can_edit_deal(u, o): return await cq.answer("Нельзя редактировать", show_alert=True)
    rows, line = [], []
    for key, label, _ in DEAL_FIELDS:
        line.append(btn(label, f"deal:ef:{did}:{key}"))
        if len(line) == 2: rows.append(line); line = []
    if line: rows.append(line)
    rows.append([btn("⬅️ Назад", f"deal:view:{did}")])
    await safe_edit(cq, "✏️ Что изменить?", kb(rows)); await cq.answer()

@dp.callback_query(F.data.startswith("deal:ef:"))
async def deal_edit_field(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id); _, _, did, key = cq.data.split(":"); did = int(did); o = await get_deal(did)
    if not can_edit_deal(u, o): return await cq.answer("Нельзя редактировать", show_alert=True)
    label = next((l for k, l, _ in DEAL_FIELDS if k == key), key)
    await state.clear(); await state.update_data(deal_id=did, field=key); await state.set_state(EditField.value)
    await cq.message.answer(f"Новое значение — <b>{esc(label)}</b>:"); await cq.answer()

@dp.message(EditField.value)
async def edit_field_save(m: Message, state: FSMContext):
    d = await state.get_data(); u = await get_user(m.from_user.id); o = await get_deal(d["deal_id"])
    if not can_edit_deal(u, o): return await state.clear()
    key = d["field"]; ftype = next((t for k, _, t in DEAL_FIELDS if k == key), "text")
    if ftype == "date":
        val = parse_date(m.text)
        if val is None: return await m.answer("Формат ДД.ММ.ГГГГ:")
    elif ftype == "int": val = parse_int(m.text)
    elif ftype == "num": val = parse_num(m.text)
    else: val = None if m.text.strip() in ("-", "") else m.text.strip()
    async with pool.acquire() as conn: await conn.execute(f"UPDATE orders SET {key}=$1 WHERE id=$2", val, d["deal_id"])
    asyncio.create_task(sync_deal_to_bitrix(d["deal_id"]))
    await state.clear(); text, markup = await render_deal(await get_deal(d["deal_id"]), u); await m.answer("✅ Обновлено.\n\n" + text, reply_markup=markup)

@dp.callback_query(F.data.startswith("deal:rm:"))
async def deal_rm(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); did = int(cq.data.split(":")[2]); o = await get_deal(did)
    if not can_edit_deal(u, o): return await cq.answer("Удалять нельзя", show_alert=True)
    await safe_edit(cq, f"⚠️ Удалить заявку #{did}?", kb([[btn("❌ Да", f"deal:rmok:{did}")], [btn("Отмена", f"deal:view:{did}")]])); await cq.answer()

@dp.callback_query(F.data.startswith("deal:rmok:"))
async def deal_rmok(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); did = int(cq.data.split(":")[2]); o = await get_deal(did)
    if not can_edit_deal(u, o): return await cq.answer("Нельзя", show_alert=True)
    if o.get("bitrix_id"):
        try:
            import httpx
            async with httpx.AsyncClient() as client:
                await client.post(f"{BITRIX_URL}crm.deal.delete.json", json={
                    "id": o["bitrix_id"],
                    "ID": o["bitrix_id"]
                })
        except Exception:
            pass
    async with pool.acquire() as conn: await conn.execute("DELETE FROM orders WHERE id=$1", did)
    await safe_edit(cq, "🗑 Удалена.", kb([[btn("⬅️ К компании", f"co:view:{o['company_id']}")]])); await cq.answer()

@dp.callback_query(F.data == "rep:menu")
async def rep_menu(cq: CallbackQuery, state: FSMContext):
    await state.clear(); u = await get_user(cq.from_user.id)
    if not is_allowed(u): return await cq.answer("Нет доступа")
    if is_admin(u):
        rows = [[btn("🌐 Общий (все)", "rep:p:all:0")], [btn("👤 По менеджеру", "rep:list:mgr")], [btn("🏢 По компании", "rep:list:comp")], [btn("🏠 Меню", "menu:main")]]
        await safe_edit(cq, "📊 Какой отчёт?", kb(rows))
    else: await safe_edit(cq, "📊 Отчёт по вашим компаниям:", period_kb("my", u["id"], "menu:main"))
    await cq.answer()

@dp.callback_query(F.data.startswith("rep:list:"))
async def rep_list(cq: CallbackQuery):
    u = await get_user(cq.from_user.id)
    if not is_admin(u): return await cq.answer("Нет доступа")
    kind = cq.data.split(":")[2]
    async with pool.acquire() as conn:
        if kind == "mgr":
            items = await conn.fetch("SELECT id,name,telegram_id FROM managers ORDER BY name")
            rows = [[btn(f"👤 {x['name'] or x['telegram_id']}", f"rep:p:mgr:{x['id']}")] for x in items]
        else:
            items = await conn.fetch("SELECT id,name FROM companies ORDER BY name")
            rows = [[btn(f"🏢 {x['name']}", f"rep:p:comp:{x['id']}")] for x in items[:80]]
    rows.append([btn("⬅️ Назад", "rep:menu")]); await safe_edit(cq, "Выберите объект:", kb(rows)); await cq.answer()

@dp.callback_query(F.data.startswith("rep:p:"))
async def rep_pick_period(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); _, _, scope, sid = cq.data.split(":"); sid = int(sid)
    if scope in ("all", "mgr") and not is_admin(u): return await cq.answer("Нет доступа")
    if scope == "comp" and not can_access_company(u, await get_company(sid)): return await cq.answer("Нет доступа")
    await safe_edit(cq, "Период отчёта:", period_kb(scope, sid, "rep:menu")); await cq.answer()

@dp.callback_query(F.data.startswith("rep:day:"))
async def rep_day(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id); _, _, scope, sid = cq.data.split(":")
    await state.clear(); await state.update_data(scope=scope, sid=int(sid)); await state.set_state(ReportDay.day)
    await cq.message.answer("📅 За какой день отчёт? (ДД.ММ.ГГГГ):"); await cq.answer()

@dp.message(ReportDay.day)
async def rep_day_go(m: Message, state: FSMContext):
    d = await state.get_data(); u = await get_user(m.from_user.id); day = parse_date(m.text)
    await state.clear()
    if day is None: return await m.answer("Неверная дата.")
    scope, sid = d["scope"], d["sid"]
    if scope == "my": sid = u["id"]
    rows = await fetch_orders(scope, sid, day, day)
    await deliver_report(m, u, rows, scope, sid, f"День {fmt_date(day)}")

@dp.callback_query(F.data.startswith("rep:go:"))
async def rep_go(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); _, _, scope, sid, period = cq.data.split(":"); sid = int(sid)
    if scope in ("all", "mgr") and not is_admin(u): return await cq.answer("Нет доступа")
    if scope == "comp" and not can_access_company(u, await get_company(sid)): return await cq.answer("Нет доступа")
    if scope == "my": sid = u["id"]
    await cq.answer("Формирую…")
    d_from, d_to = period_range(period)
    rows = await fetch_orders(scope, sid, d_from, d_to)
    await deliver_report(cq.message, u, rows, scope, sid, PERIODS[period][0])

async def deliver_report(target, u, rows, scope, sid, plabel):
    if scope == "all": title = "Общий отчёт"
    elif scope == "comp": title = f"Отчёт: {(await get_company(sid))['name']}"
    elif scope == "mgr":
        async with pool.acquire() as conn: mn = await conn.fetchrow("SELECT name,telegram_id FROM managers WHERE id=$1", sid)
        title = f"Менеджер: {mn['name'] or mn['telegram_id']}"
    else: title = "Отчёт по моим компаниям"
    if not rows: return await target.answer(f"За период «{plabel}» данных нет.")
    path = build_excel(rows, title, plabel)
    await target.answer_document(FSInputFile(path), caption=f"📊 {title}\nПериод: {plabel}\nЗаявок: {len(rows)}")

@dp.callback_query(F.data == "adm:auth")
async def adm_auth(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id)
    if not is_admin(u): return await cq.answer("Только для руководителя", show_alert=True)
    await state.set_state(AdminAuth.password); await cq.message.answer("🔐 Пароль администратора:"); await cq.answer()

@dp.message(AdminAuth.password)
async def adm_check(m: Message, state: FSMContext):
    await state.clear(); u = await get_user(m.from_user.id)
    if not is_admin(u): return await m.answer("Доступ запрещён.")
    if m.text.strip() != ADMIN_PASSWORD: return await m.answer("❌ Неверный пароль.", reply_markup=main_menu(u))
    await m.answer("✅ Админка открыта.", reply_markup=kb([[btn("👥 Менеджеры", "adm:mgrs")], [btn("📊 Отчёты", "rep:menu")], [btn("🚪 Выйти", "menu:main")]]))

@dp.callback_query(F.data == "adm:mgrs")
async def adm_mgrs(cq: CallbackQuery):
    u = await get_user(cq.from_user.id)
    if not is_admin(u): return await cq.answer("Нет доступа")
    async with pool.acquire() as conn: ms = await conn.fetch("SELECT * FROM managers ORDER BY role DESC, name")
    rows = []
    for x in ms:
        flag = "👑" if x["role"] == "admin" else ("🟢" if x["active"] else "🔴")
        rows.append([btn(f"{flag} {x['name'] or x['telegram_id']}", f"adm:mgr:{x['id']}")])
    rows.append([btn("➕ Добавить менеджера", "adm:add"), btn("⬅️ Назад", "adm:back")])
    await safe_edit(cq, "👥 <b>Менеджеры</b>", kb(rows)); await cq.answer()

@dp.callback_query(F.data == "adm:back")
async def adm_back(cq: CallbackQuery): await safe_edit(cq, "Админка:", kb([[btn("👥 Менеджеры", "adm:mgrs")], [btn("📊 Отчёты", "rep:menu")], [btn("🚪 Выйти", "menu:main")]])); await cq.answer()

@dp.callback_query(F.data.startswith("adm:mgr:"))
async def adm_mgr_view(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); mid = int(cq.data.split(":")[2])
    if not is_admin(u): return await cq.answer("Нет доступа")
    async with pool.acquire() as conn:
        x = await conn.fetchrow("SELECT * FROM managers WHERE id=$1", mid)
        nco = (await conn.fetchval("SELECT COUNT(*) FROM companies WHERE manager_id=$1", mid))
        admins = (await conn.fetchval("SELECT COUNT(*) FROM managers WHERE role='admin' AND active=1"))
    if not x: return await cq.answer("Не найдено")
    txt = f"👤 <b>{esc(x['name'])}</b>\nID: <code>{x['telegram_id']}</code>\nРоль: {x['role']} · {'активен' if x['active'] else 'отключён'}\nКомпаний: {nco}"
    rows = []
    if x["role"] == "admin":
        if admins > 1: rows.append([btn("🙍 Сделать менеджером", f"adm:role:{mid}:manager")])
    else:
        rows.append([btn("👑 Сделать руководителем", f"adm:role:{mid}:admin")])
        rows.append([btn("🔴 Отключить" if x["active"] else "🟢 Включить", f"adm:toggle:{mid}:{0 if x['active'] else 1}")])
        rows.append([btn("✏️ Переименовать", f"adm:ren:{mid}")])
    rows.append([btn("📊 Отчёт", f"rep:p:mgr:{mid}"), btn("⬅️ К списку", "adm:mgrs")])
    await safe_edit(cq, txt, kb(rows)); await cq.answer()

@dp.callback_query(F.data.startswith("adm:role:"))
async def adm_role(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); _, _, mid, role = cq.data.split(":")
    if not is_admin(u): return await cq.answer("Нет доступа")
    async with pool.acquire() as conn: await conn.execute("UPDATE managers SET role=$1, active=1 WHERE id=$2", role, int(mid))
    cq.data = f"adm:mgr:{mid}"; await adm_mgr_view(cq)

@dp.callback_query(F.data.startswith("adm:toggle:"))
async def adm_toggle(cq: CallbackQuery):
    u = await get_user(cq.from_user.id); _, _, mid, val = cq.data.split(":")
    if not is_admin(u): return await cq.answer("Нет доступа")
    async with pool.acquire() as conn: await conn.execute("UPDATE managers SET active=$1 WHERE id=$2 AND role!='admin'", int(val), int(mid))
    cq.data = f"adm:mgr:{mid}"; await adm_mgr_view(cq)

@dp.callback_query(F.data.startswith("adm:ren:"))
async def adm_ren(cq: CallbackQuery, state: FSMContext):
    u = await get_user(cq.from_user.id)
    if not is_admin(u): return await cq.answer("Нет доступа")
    await state.set_state(RenameManager.name); await state.update_data(mid=int(cq.data.split(":")[2]))
    await cq.message.answer("Новое имя:"); await cq.answer()

@dp.message(RenameManager.name)
async def adm_ren_save(m: Message, state: FSMContext):
    u = await get_user(m.from_user.id)
    if not is_admin(u): return await state.clear()
    d = await state.get_data(); await state.clear()
    async with pool.acquire() as conn: await conn.execute("UPDATE managers SET name=$1 WHERE id=$2", m.text.strip(), d["mid"])
    await m.answer("✅ Имя обновлено.", reply_markup=kb([[btn("👥 К менеджерам", "adm:mgrs")]]))

@dp.callback_query(F.data == "adm:add")
async def adm_add(cq: CallbackQuery, state: FSMContext):
    if not is_admin(await get_user(cq.from_user.id)): return await cq.answer("Нет доступа")
    await state.set_state(ManagerForm.telegram_id); await cq.message.answer("➕ Telegram ID:"); await cq.answer()

@dp.message(ManagerForm.telegram_id)
async def adm_add_id(m: Message, state: FSMContext):
    if not is_admin(await get_user(m.from_user.id)): return await state.clear()
    digits = "".join(ch for ch in m.text if ch.isdigit())
    if not digits: return await m.answer("Нужно число:")
    await state.update_data(tid=int(digits)); await state.set_state(ManagerForm.name); await m.answer("Имя менеджера:")

@dp.message(ManagerForm.name)
async def adm_add_name(m: Message, state: FSMContext):
    u = await get_user(m.from_user.id)
    if not is_admin(u): return await state.clear()
    d = await state.get_data(); await state.clear()
    try:
        async with pool.acquire() as conn: await conn.execute("INSERT INTO managers(telegram_id,name,role,active,created_at) VALUES($1,$2,'manager',1,$3)", d["tid"], m.text.strip(), datetime.now().isoformat())
        await m.answer(f"✅ Добавлен.", reply_markup=kb([[btn("👥 К менеджерам", "adm:mgrs")]]))
    except asyncpg.exceptions.UniqueViolationError:
        await m.answer("Этот ID уже есть.", reply_markup=kb([[btn("👥 К менеджерам", "adm:mgrs")]]))

@dp.message(CommandStart())
@dp.message(Command("start"))
@dp.message(Command("desktop"))
async def cmd_start(m: Message):
    u = await get_user(m.from_user.id)
    if not is_allowed(u): return await m.answer(f"⛔️ Доступ запрещён. ID: <code>{m.from_user.id}</code>")
    
    web_url = os.environ.get("WEB_APP_URL", "https://kazoc-tgbot-nine.vercel.app")
    desktop_url = f"{web_url}/?tg_id={m.from_user.id}"
    
    text = (
        f"👋 <b>Добро пожаловать в KazOC CRM!</b>\n\n"
        f"🖥 <b>Ссылка для ПК (Десктопный режим):</b>\n"
        f"<code>{desktop_url}</code>\n\n"
        f"<i>Скопируйте эту ссылку и откройте её на вашем компьютере, чтобы работать в удобном широком окне с двумя колонками!</i>"
    )
    
    kb_rows = [
        [InlineKeyboardButton(text="📱 Открыть Web App (в TG)", web_app=WebAppInfo(url=web_url))],
        [InlineKeyboardButton(text="🖥 Открыть на ПК", url=desktop_url)],
        [btn("📋 Компании-клиенты", "co:list:0")],
        [btn("➕ Новая компания", "co:new")],
        [btn("📊 Отчёты", "rep:menu")]
    ]
    if is_admin(u):
        kb_rows.append([btn("🔐 Админка", "adm:auth")])
        
    await m.answer(text, reply_markup=kb(kb_rows))

@dp.message(StateFilter(None))
async def fallback(m: Message):
    u = await get_user(m.from_user.id)
    if not is_allowed(u): return await m.answer(f"⛔️ Доступ запрещён. ID: <code>{m.from_user.id}</code>")
    await m.answer("Меню:", reply_markup=main_menu(u))

async def startup():
    global pool
    pool = await asyncpg.create_pool(DATABASE_URL, statement_cache_size=0)
    await init_db()

async def main():
    await startup()
    bot = Bot(token=BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    print("Бот запущен.")
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
